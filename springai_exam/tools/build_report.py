# -*- coding: utf-8 -*-
"""生成《解轲尧20233221作业二.docx》。

版式对齐作业一的报告：标题 + 姓名/学号/班级 + 「N、任务X」一级标题 +
「N.M 小节」二级标题 + 「描述：」「截图：」+ 虚线截图位。

截图位里不插图，改写成「打开哪儿 / 截什么」的操作说明，由用户截图后自己粘进去。
"""
import os

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(os.path.dirname(ROOT), "解轲尧20233221作业二.docx")

GRAY = (0x8C, 0x8C, 0x8C)
BLUE = (0x1F, 0x4E, 0x79)

doc = Document()

# 全局默认字体
normal = doc.styles["Normal"]
normal.font.name = "宋体"
normal.font.size = Pt(11)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _setfont(run, size, bold=False, font="宋体", color=None):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def para(text, size=11, bold=False, font="宋体", align=None, space_after=4,
         color=None, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.25
    if align is not None:
        p.alignment = align
    if indent is not None:
        p.paragraph_format.left_indent = Inches(indent)
    _setfont(p.add_run(text), size, bold, font, color)
    return p


def h1(text):
    para(text, size=14, bold=True, font="黑体", space_after=6)


def h2(text):
    para(text, size=12, bold=True, font="黑体", space_after=4)


def _dashed_borders(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "dashed")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:color"), "BFBFBF")
        borders.append(el)
    table._tbl.tblPr.append(borders)


def shot(steps, height_cm=5.4):
    """虚线截图位：第 1 行提示粘贴，后面几行写「打开哪儿 / 截什么」。"""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    _dashed_borders(t)
    row = t.rows[0]
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    row.height = Cm(height_cm)

    cell = t.cell(0, 0)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    _setfont(p.add_run("【 截图位 · 截好后把图片直接插入到这个框里 】"), 10, True, color=GRAY)

    for line in steps:
        q = cell.add_paragraph()
        q.paragraph_format.space_after = Pt(2)
        q.paragraph_format.line_spacing = 1.2
        _setfont(q.add_run(line), 10.5, color=BLUE)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def desc(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.25
    _setfont(p.add_run("描述："), 11, True)
    _setfont(p.add_run(text), 11)
    return p


def label(text="截图："):
    para(text, size=11, space_after=2)


# ============================ 封面 ============================
para("作业【校园百事通 —— Spring AI 领域智能助手】", size=16, bold=True,
     font="黑体", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
para("姓名：解轲尧　　学号：20233221　　班级：软件231", size=11,
     align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)

# ============================ 领域与角色 ============================
h1("0、所选领域与角色设定")
desc("领域：校园学习与校园生活。助手名叫「小帮」，身份是「校园百事通」，只解答选课、考试、"
     "图书馆、食堂、宿舍、社团、奖学金、实习、考研这一类校园相关的问题；问代码、写小说、"
     "医疗诊断等校园之外的问题时会礼貌说明超出服务范围并引导回校园话题。")
desc("角色规则一共 5 条（要求不少于 3 条）：① 只回答校园范围内的问题；② 越界问题礼貌拒绝、"
     "不强行作答；③ 一律用 Markdown 输出，先给结论再分点；④ 建议要具体可执行并结合校园场景；"
     "⑤ 每次回答结尾固定加上「以上回答希望对你有帮助！」。")
desc("这份人设不写在 Java 里，而是放在 resources/prompt/exam_prompt.txt，接口通过 "
     "@Value(\"classpath:prompt/exam_prompt.txt\") Resource 读取后交给 .system(...)，"
     "改人设只要改 txt，不用动代码。")
label()
shot([
    "① 打开：IDEA 左侧 Project 面板 → src/main/resources/prompt/exam_prompt.txt（双击打开）。",
    "② 截取：整个文件的内容，要能看清开头的【系统角色与身份】和下面的 5 条规则、以及第 5 条的固定结尾语。",
])

# ============================ 任务一 ============================
h1("1、任务一：环境与项目搭建")

h2("1.1 项目与依赖")
desc("新建 Spring Boot 4.1.0 项目，JDK 17（本机 17.0.19），Maven 3.9.14。项目名 springai_exam，"
     "包名 com.example.exam。pom.xml 用 spring-ai-bom 2.0.0 统一管理 Spring AI 版本，"
     "引入 spring-ai-starter-model-openai 与 lombok。")
desc("需要说明一处：作业文档里写的依赖名 spring-ai-starter-model-webmvc 在 Maven 中央仓库"
     "并不存在（已核对，检索结果为 0 条），实际能拉到的 Web 依赖是 spring-boot-starter-webmvc"
     "（Spring Boot 4 里 spring-boot-starter-web 的改名），本作业用的就是后者。")
label()
shot([
    "① 打开：IDEA 里的 pom.xml。",
    "② 截取：<parent>（spring-boot-starter-parent 4.1.0）、spring-ai.version 2.0.0、",
    "      spring-ai-bom 的 <dependencyManagement> 片段，以及 <dependencies> 里的 spring-ai-starter-model-openai。",
    "　　文件太长的话分两张截，上下拼在一起即可。",
])

h2("1.2 application.yaml 配置")
desc("端口 8080；API Key 从环境变量读取，不写死在代码里；base-url 用兼容 OpenAI 协议的服务地址；"
     "对话模型 deepseek-chat；开启 spring.servlet.encoding.force 强制 UTF-8 解决中文乱码；"
     "并把 org.springframework.ai 的日志级别调到 debug。")
desc("本机配置的环境变量是 DEEPSEEK_API_KEY，所以 base-url 指向 DeepSeek 的兼容地址。"
     "换成阿里云百炼只需改两行：api-key 改成 ${DASHSCOPE_API_KEY:...}，"
     "base-url 改成 https://dashscope.aliyuncs.com/compatible-mode/v1，"
     "model 改成百炼的模型名，Java 代码一行都不用动。")
label()
shot([
    "① 打开：IDEA 里的 src/main/resources/application.yaml。",
    "② 截取：整个文件（要能看清 server.port 8080、${DEEPSEEK_API_KEY}、base-url、chat.model、",
    "      servlet.encoding.force 和 logging.level 这几处）。",
])

h2("1.3 项目启动成功")
desc("控制台出现 Tomcat started on port 8080 与 Started SpringaiExamApplication，"
     "启动耗时 3.09 秒，无报错。")
label()
shot([
    "① 在 IDEA 里点绿色三角运行 SpringaiExamApplication。",
    "　　⚠️ 启动前先确认 8080 没被占：之前命令行起的那个实例要先关掉，否则会报 Port 8080 was already in use。",
    "② 截取：Run 控制台最后十来行，要能看到 Tomcat started on port 8080 和",
    "      Started SpringaiExamApplication in x.xx seconds 这两句。",
])

# ============================ 任务二 ============================
h1("2、任务二：同步对话接口")
h2("2.1 接口实现")
desc("GET /exam/chat，参数 question（String，默认值「你是谁」），返回 String。"
     "实现就是 chatClient.prompt().user(question).call().content()，用 call() 等模型把回答"
     "全部生成后一次性返回。四个任务的接口都写在 ExamController 里：")
label("代码：")
shot([
    "① 打开：IDEA 里的 src/main/java/com/example/exam/controller/ExamController.java。",
    "② 截取：整个类（四个任务的接口方法都在这一个文件里，建议整屏截）。",
])

h2("2.2 传参访问")
desc("访问 http://localhost:8080/exam/chat?question=你好，返回一段完整回答，耗时 4 秒左右。"
     "（因为全局挂了自定义 Advisor，回答会按「先给结论、再分点」的结构输出，这是任务五的预期效果。）")
label()
shot([
    "① 打开：浏览器地址栏访问 http://localhost:8080/exam/chat?question=你好",
    "② 截取：返回的回答全文，地址栏里的网址要一起截进去。",
    "　　💡 这个接口返回的是纯字符串，浏览器按 HTML 渲染会把换行压掉、Markdown 的 ** 原样显示。",
    "　　　 想看带换行的原始返回，可按 F12 → Network → 点那条 chat 请求 → 右侧 Response 标签再截。",
])

h2("2.3 不传参数访问")
desc("访问 http://localhost:8080/exam/chat，没有传 question，接口按默认值「你是谁」提问，"
     "正常返回模型对自身身份的说明。")
label()
shot([
    "① 打开：浏览器地址栏访问 http://localhost:8080/exam/chat（注意网址里不带 question 参数）。",
    "② 截取：返回的回答，以及地址栏 —— 要能看出这次没有传参数，走的是默认问题。",
])

# ============================ 任务三 ============================
h1("3、任务三：流式对话 + 角色预设")

h2("3.1 流式接口实现")
desc("GET /exam/stream，参数 question（String，默认值「你能做什么」），返回 Flux<String>，"
     "把 call() 换成 stream() 即可。底层是 SSE / chunked 分块传输，文字是逐块推给浏览器的。"
     "接口代码见 2.1 的 ExamController 截图。")

h2("3.2 流式效果")
desc("下面是抓到的真实响应：同一个回答被切成了 20 多块，每块到达的时间都单独记录下来，"
     "从 +748ms 一直到 +2900ms 分散在不同时刻，说明不是一次性返回，而是逐块实时到达。")
label()
shot([
    "① 打开：浏览器地址栏访问 http://localhost:8080/exam/stream?question=图书馆几点开门",
    "② 截取：★ 趁回答还在一个字一个字往外冒的时候按截图键 ★",
    "　　—— 图里要能同时看到「只出来一半的文字」+「标签页还在转圈 / 状态栏还在加载」，",
    "　　　 这才看得出是逐块返回；等它全出完再截就和一次性返回没区别了。",
    "　　💡 更严谨的做法：F12 → Network → 点那条 stream 请求 → Response，能看到内容一段段添上去。",
    "　　💡 对比着看更明显：/exam/chat 是等好几秒后整段突然出现，这个是一点点往外冒。",
])

h2("3.3 越界提问被礼貌拒绝")
desc("用「帮我写一段冒泡排序的Java代码」这种明显超出校园范围的问题去问，助手先说明"
     "「这超出了我的服务范围」，表明自己只负责校园相关的问题，再把话题引回选课、考试、图书馆等"
     "校园方向，全程没有强行作答；结尾依然带着固定结束语。")
label()
shot([
    "① 打开：浏览器地址栏访问",
    "　　http://localhost:8080/exam/stream?question=帮我写一段冒泡排序的Java代码",
    "② 截取：回答里「这超出了我的服务范围」那段说明，以及最后那句固定结束语。",
])

h2("3.4 全局跨域配置")
desc("新建 config/CorsConfig.java 实现 WebMvcConfigurer，对 /** 放开 allowedOriginPatterns(\"*\")、"
     "允许 GET/POST/PUT/DELETE/OPTIONS、允许所有请求头、允许携带凭证，预检请求缓存 1 小时。"
     "一次配置全局生效，不用在每个 Controller 上重复加 @CrossOrigin。")
desc("实测：模拟前端跑在 http://localhost:5173 发起跨域请求，响应头带上了 "
     "Access-Control-Allow-Origin: http://localhost:5173；对应的 OPTIONS 预检请求也返回 200，"
     "并正确给出 Allow-Methods / Allow-Headers / Max-Age。")
label()
shot([
    "⚠️ 直接在地址栏访问是看不出跨域的（同一个页面请求自己不算跨域，两端都不会加跨域头），",
    "　　 必须造一个真正跨域的请求。下面这套操作不需要额外启前端项目：",
    "① 浏览器先打开 http://localhost:8080/exam/chat?question=你好 当跳板页。",
    "② 按 F12 → Console，粘贴下面这行回车：",
    "　　fetch('http://127.0.0.1:8080/exam/chat?question=跨域测试',{headers:{'X-Test':'1'}})",
    "　　  .then(r=>r.text()).then(t=>console.log('跨域成功：',t))",
    "　　（localhost 和 127.0.0.1 在浏览器眼里是两个不同的域，所以这是真跨域；",
    "　　　带上 X-Test 是为了让浏览器先发一条 OPTIONS 预检请求。）",
    "③ 截取：Network 面板里新出现的两条请求（一条 OPTIONS 预检、一条 GET），点开 Headers 标签，",
    "　　 要能看到 access-control-allow-origin / allow-methods / max-age 这几个响应头。",
    "　　✅ 已实测：服务端这几项都会正常返回，Console 不会报 CORS 错误。",
], height_cm=7.0)

# ============================ 任务四 ============================
h1("4、任务四：提示词工程")

h2("4.1 提示词外部化")
desc("任务三的角色提示词已经抽到 src/main/resources/prompt/exam_prompt.txt，"
     "控制器里用 @Value(\"classpath:prompt/exam_prompt.txt\") 注入成 Resource，再传给 .system(...)。"
     "实测：把 txt 里第 5 条规则的结尾语改成「这条结尾语是从 exam_prompt.txt 读出来的，改文件就生效！」，"
     "不碰任何 Java 代码直接重启，同一个问题的回答结尾就跟着变了 —— 人设确实由外部文件决定。")
desc("这里有一个实测踩到的坑：.system(Resource) 的单参重载内部走的是 Charset.defaultCharset()，"
     "而中文 Windows 上 JDK 17 的默认字符集是 GBK，会把 UTF-8 的 txt 读成乱码"
     "（控制台能看到 SystemMessage 的内容变成「銆愮郴缁熻鑹蹭笌韬唤」）。"
     "所以这里改用了带字符集参数的重载 .system(examPrompt, StandardCharsets.UTF_8) "
     "显式指定 UTF-8，问题解决。")
label()
shot([
    "① 打开：IDEA 里的 src/main/resources/prompt/exam_prompt.txt，",
    "　　 把第 5 条规则结尾的「以上回答希望对你有帮助！」改成别的句子，",
    "　　 例如「这条结尾语是从 exam_prompt.txt 读出来的，改文件就生效！」—— 不要动任何 Java 代码。",
    "② 重启项目，浏览器再访问同一个 http://localhost:8080/exam/chat?question=你好",
    "③ 截取：回答的最后一句已经变成你刚改的新句子（能同时带上改过的 txt 画面更好）。",
    "　　⚠️ 截完记得把 txt 改回原样，别把答案留在提交版本里。",
], height_cm=6.2)

h2("4.2 PromptTemplate 动态拼接")
desc("GET /exam/template，参数 question（用户问题）和 topic（主题），返回 String，用的是 call()。"
     "模板里有两个占位符 {topic} 和 {question}，用 PromptTemplate.builder().template(...).build()"
     ".create(Map.of(\"topic\", topic, \"question\", question)) 完成替换。")
desc("访问 http://localhost:8080/exam/template?topic=食堂&question=中午吃什么 ，"
     "回答里同时体现了「食堂」这个主题和「中午吃什么」这个具体问题。")
label()
shot([
    "① 打开：浏览器地址栏访问",
    "　　http://localhost:8080/exam/template?topic=食堂&question=中午吃什么",
    "　　（网址里的中文直接粘进浏览器即可，浏览器会自己转码。）",
    "② 截取：返回的回答，要能看出它同时用到了「食堂」这个主题和「中午吃什么」这个问题。",
])

# ============================ 任务五 ============================
h1("5、任务五：Advisor 拦截链")

h2("5.1 三个 Advisor 的挂载")
desc("在 AiConfig 里用 .defaultAdvisors(...) 一次挂上三个拦截器："
     "SimpleLoggerAdvisor（日志）、SafeGuardAdvisor（敏感词）、RewritePromptAdvisor（自定义）。"
     "执行顺序由各自的 getOrder() 决定，数值越小越先执行。")
desc("SimpleLoggerAdvisor 是框架默认的 order 0，最先打印请求；SafeGuardAdvisor 显式传 1，"
     "排在第二个，命中敏感词就直接返回兜底语、不再往下走；自定义的 RewritePromptAdvisor 用 2，"
     "排在 SafeGuardAdvisor 之后，所以敏感词一定是在原文上被拦下来的，不会被改写干扰。"
     "请求阶段按 order 从小到大依次进入 before，响应阶段反过来返回，是个洋葱模型。")
label()
shot([
    "① 打开：IDEA 里的 src/main/java/com/example/exam/config/AiConfig.java。",
    "② 截取：整个类，重点要看清 .defaultAdvisors(...) 里的三个 Advisor 和上面那串敏感词列表。",
])

h2("5.2 日志 Advisor（SimpleLoggerAdvisor）")
desc("yaml 里把 logging.level.org.springframework.ai 设为 debug 之后，控制台能看到每次对话的"
     "完整请求与响应。下面摘的是真实抓到的四组日志，按发生顺序排列：")
para("① 第一组是 /exam/chat?question=你好 的请求与响应，能看到 content 里就是用户原话，"
     "响应之间隔了约 3.9 秒（模型真实生成耗时）。", indent=0.2)
para("② 第二行是 /exam/stream 的请求，messages 里能明显看到 SystemMessage{textContent='【系统角色与身份】…'}，"
     "说明外部化到 txt 的人设确实被拼进了提示词。", indent=0.2)
para("③ 第三组是 /exam/template 的请求与响应，请求里带着渲染好的「请围绕「食堂」这一主题…」，"
     "间隔同样在秒级。", indent=0.2)
para("④ 最后一组是「帮我代写一篇论文」的请求与响应 —— 两条日志的时间戳只差 2 毫秒"
     "（20.555 → 20.557），对比前面几组动辄三四秒，可以直接断定这次没有请求模型，"
     "是 SafeGuardAdvisor 命中敏感词后直接返回了兜底语。", indent=0.2)
label()
shot([
    "① 先在浏览器访问一次 http://localhost:8080/exam/chat?question=你好，让控制台产生日志。",
    "② 回 IDEA 的 Run 控制台，按 Ctrl+F 搜 SimpleLoggerAdvisor 定位，",
    "　　 截取一条「请求」日志 + 一条「响应」日志，时间戳要露出来（能看出中间隔了几秒）。",
    "③ 再访问 http://localhost:8080/exam/chat?question=帮我代写一篇论文，回控制台截这两条 ——",
    "　　 它们的时间戳只差几毫秒，和上面隔好几秒的一对比，就能证明这次没请求大模型。",
    "　　⚠️ 控制台日志很长，先搜关键字定位再截，别截一整屏看不清重点。",
], height_cm=6.2)

h2("5.3 敏感词 Advisor（SafeGuardAdvisor）")
desc("敏感词列表自定义了 7 个（要求不少于 5 个）：代写、刷单、赌博、色情、暴力、诈骗、毒品。"
     "选的都和校园场景有关，正常提问不会误伤。命中后直接返回兜底语，不会请求大模型 ——"
     "下面这次「帮我代写一篇论文」的耗时是 0.00 秒，对比其它接口动辄 2～5 秒，"
     "可以直接看出它没有走网络请求。")
label()
shot([
    "① 打开：浏览器地址栏访问 http://localhost:8080/exam/chat?question=帮我代写一篇论文",
    "② 截取：返回的是兜底语（「抱歉，您的问题里包含不适合讨论的内容……」，不是模型的回答）。",
    "③ 想加一份「没走模型」的证据，可以再截一张 IDEA 控制台里那条请求 / 响应只差 2 毫秒的日志。",
])

h2("5.4 自定义 Advisor（RewritePromptAdvisor）")
desc("实现 BaseAdvisor 接口，重写 before / after / getOrder 三个方法。"
     "before(...) 里用 PromptTemplate 把用户的原始问题包进一个「先给结论、再分点解释、"
     "Markdown 排版」的提问框架，然后用新的 Prompt 替换原请求；after(...) 原样返回；"
     "getOrder() 返回 2，排在 SafeGuardAdvisor（1）之后。")
desc("有一处对讲义示例做了修改：讲义是把整个 Prompt 换掉，实测那样会连同 .system(...) 设好的"
     "人设一起丢掉，校园百事通的角色和越界拒绝都会失效。所以这里只替换「用户消息」的内容，"
     "System 消息保留不动，只重写用户问题 —— 这也更贴合作业里「改写用户问题」的表述。")
label()
shot([
    "① 打开：IDEA 里的 src/main/java/com/example/exam/config/RewritePromptAdvisor.java。",
    "② 截取：整个类，before / after / getOrder 三个方法都要在画面里。",
])

h2("5.5 挂载前后的对比")
desc("用的是同一个问题「什么是多态」。挂载前是模型的自由发挥：开头是「多态是面向对象编程的"
     "三大核心特性之一（另外两个是封装和继承）。简单来说，多态的意思是……」，结构由模型自己决定。")
label()
shot([
    "① 打开 IDEA 里的 config/AiConfig.java，把 new RewritePromptAdvisor() 这一行用 // 注释掉。",
    "　　⚠️ 它原来是 .defaultAdvisors(...) 里的最后一项，注释掉之后要把**上一行末尾的逗号删掉**，",
    "　　　 Java 的方法参数列表不允许尾逗号，不改会编译报错（这一步我踩过）。",
    "② 重启项目，浏览器访问 http://localhost:8080/exam/chat?question=什么是多态",
    "③ 截取：这时是模型的自由发挥 —— 开头是一整段话，结构由模型自己定。",
])

desc("挂载后，同一个问题的回答第一句变成了加粗的一句话结论"
     "「多态就是“同一个行为，在不同对象上表现出不同结果”的能力。」，"
     "后面严格按 1. 2. 3. 分点展开 —— 输入提示词被改写这件事在输出上是可以直接看出来的。")
label()
shot([
    "① 把上一步注释掉的 new RewritePromptAdvisor() 取消注释（并把逗号加回去），重启项目。",
    "② 同一个问题：浏览器访问 http://localhost:8080/exam/chat?question=什么是多态",
    "③ 截取：回答的第一句是被改写成加粗的一句话结论，后面严格按 1. 2. 3. 分点展开。",
    "　　（两张放一起对比，格式变化一目了然。）",
])

# ============================ 接口汇总 ============================
h1("6、接口测试 URL 汇总")
urls = [
    ("任务二", "同步对话", "http://localhost:8080/exam/chat?question=你好"),
    ("任务二", "默认问题", "http://localhost:8080/exam/chat"),
    ("任务三", "流式对话", "http://localhost:8080/exam/stream?question=图书馆几点开门"),
    ("任务三", "越界拒绝", "http://localhost:8080/exam/stream?question=帮我写一段冒泡排序的Java代码"),
    ("任务四", "模板拼接", "http://localhost:8080/exam/template?topic=食堂&question=中午吃什么"),
    ("任务五", "敏感词兜底", "http://localhost:8080/exam/chat?question=帮我代写一篇论文"),
]
t = doc.add_table(rows=1, cols=3)
t.style = "Table Grid"
for i, txt in enumerate(["任务", "用途", "完整测试 URL"]):
    c = t.cell(0, i)
    c.text = ""
    _setfont(c.paragraphs[0].add_run(txt), 10.5, True)
for task, use, url in urls:
    row = t.add_row()
    for i, txt in enumerate((task, use, url)):
        row.cells[i].text = ""
        _setfont(row.cells[i].paragraphs[0].add_run(txt), 10.5)
for row in t.rows:
    row.cells[0].width = Inches(0.8)
    row.cells[1].width = Inches(1.0)
    row.cells[2].width = Inches(4.3)
doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ============================ 完成情况与问题 ============================
h1("7、完成情况与遇到的问题")

h2("7.1 完成情况")
para("任务一：已完成。项目可正常启动，pom 依赖、BOM 版本、yaml 配置、环境变量读取 Key 均按要求完成。",
     indent=0.2)
para("任务二：已完成。/exam/chat 支持传参与默认值，用 call() 返回完整字符串。", indent=0.2)
para("任务三：已完成。stream() 返回 Flux<String> 实现逐块输出；角色规则 5 条；越界提问礼貌拒绝；"
     "全局跨域已配置。", indent=0.2)
para("任务四：已完成。人设外置到 exam_prompt.txt 并用 Resource 读取；/exam/template 用 "
     "PromptTemplate 完成了 {topic} 与 {question} 两个占位符的替换。", indent=0.2)
para("任务五：已完成。日志、敏感词、自定义三个 Advisor 都已挂载，debug 日志、兜底回复、"
     "挂载前后对比三份证据齐全。", indent=0.2)

h2("7.2 遇到的问题与解决")
para("问题 1：作业文档里的依赖名 spring-ai-starter-model-webmvc 拉不到。"
     "到 Maven 中央仓库检索确认该坐标不存在（0 条结果）。"
     "解决：改用 Spring Boot 4 里对应的 spring-boot-starter-webmvc。", indent=0.2)
para("问题 2：讲义里自定义 Advisor 的 import 写的是 "
     "org.springframework.ai.chat.client.advisor.BaseAdvisor，照抄编译不过。"
     "解决：反编译 spring-ai-client-chat-2.0.0.jar 确认，BaseAdvisor 实际位于 "
     "org.springframework.ai.chat.client.advisor.api 包下，改成正确的包名即可。", indent=0.2)
para("问题 3：.system(Resource) 读 txt 出现中文乱码，控制台能看到 SystemMessage 的内容变成"
     "「銆愮郴缁熻鑹蹭笌韬唤」。原因查出来是：这个重载内部调用的是 Charset.defaultCharset()，"
     "而中文 Windows 上 JDK 17 的默认字符集是 GBK，UTF-8 的 txt 被按 GBK 解了。"
     "解决：改用带字符集参数的重载 .system(examPrompt, StandardCharsets.UTF_8)。"
     "（讲义 5-3 的常见问题 Q1 也提到过这个乱码坑，但没点明原因是默认字符集。）", indent=0.2)
para("问题 4：按讲义示例写 before() 时，把整个 Prompt 替换掉会连 System 消息一起丢掉，"
     "人设和越界拒绝全部失效。解决：只重写最后一条用户消息，用 prompt.mutate().messages(...) "
     "构造新 Prompt，System 消息原样保留。", indent=0.2)
para("问题 5：控制台日志重定向到文件后中文变成乱码。原因是 JVM 输出流的编码跟随 Windows 代码页。"
     "解决：启动时加 --logging.charset.console=UTF-8（或在 IDEA 的 VM options 里加 "
     "-Dfile.encoding=UTF-8）。", indent=0.2)

# ============================ 自检表 ============================
h1("8、测试用例自检表")
checks = [
    ("1", "启动项目", "8080 端口正常，无报错"),
    ("2", "GET /exam/chat?question=你好", "返回一段完整文字回答"),
    ("3", "不传 question 访问 /exam/chat", "按默认问题「你是谁」回答"),
    ("4", "GET /exam/stream?question=...", "文字逐字 / 逐块出现"),
    ("5", "问领域内问题", "按人设、Markdown、带固定结尾"),
    ("6", "问领域外问题", "礼貌拒绝，不强行作答"),
    ("7", "修改 exam_prompt.txt 后重启", "人设随之改变，Java 未改动"),
    ("8", "GET /exam/template?topic=&question=", "回答同时体现 topic 与 question"),
    ("9", "任意对话", "控制台打印 Spring AI debug 日志"),
    ("10", "提问中含敏感词", "直接返回兜底语，不请求模型"),
    ("11", "挂载自定义 Advisor 前后对比", "输出结构 / 格式发生预期变化"),
    ("12", "跨域请求（前端不同端口）", "不被浏览器拦截"),
]
t2 = doc.add_table(rows=1, cols=4)
t2.style = "Table Grid"
for i, txt in enumerate(["编号", "操作", "期望结果", "是否通过"]):
    c = t2.cell(0, i)
    c.text = ""
    _setfont(c.paragraphs[0].add_run(txt), 10.5, True)
for no, op, exp in checks:
    row = t2.add_row()
    for i, txt in enumerate((no, op, exp, "☑ 通过")):
        row.cells[i].text = ""
        _setfont(row.cells[i].paragraphs[0].add_run(txt), 10.5)
for row in t2.rows:
    row.cells[0].width = Inches(0.5)
    row.cells[1].width = Inches(2.2)
    row.cells[2].width = Inches(2.6)
    row.cells[3].width = Inches(0.8)

doc.save(OUT)
print("已生成:", OUT)
print("段落数:", len(doc.paragraphs), "表格数:", len(doc.tables))
