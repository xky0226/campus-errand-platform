# -*- coding: utf-8 -*-
"""生成简历：一份可编辑 .docx + 一份可直接打印/发送的 .pdf，内容同源。

一次生成三套（同一套排版代码，只换内容）：
    测试 · AI 应用开发 · Java 后端

用法: python -X utf8 tools/build_resume.py
      python -X utf8 tools/build_resume.py 测试      # 只生成某一套
"""
import os
import sys

OUT_DIR = r"C:\Users\Administrator\Desktop"
BASE = "解轲尧20233221个人求职简历"

NAME = "解轲尧"
LINE2 = "2027 届本科毕业生　|　21 岁　|　男　|　汉　|　共青团员　|　石家庄"
LINE3 = "13111566735　|　2245055027@qq.com　|　github.com/xky0226"

# 块类型: h 小节标题 / kv 项目标题(左粗+右灰) / p 正文 / b 条目(标签粗+正文)

# ══════════════════════════════════════════════════════════════════
COMMON_EDU = [
    ("h", "教育背景"),
    ("kv", "柳州工学院　软件工程（本科）", "2023.09 – 2027.06"),
    ("p", "主修课程：数据结构与算法、数据库系统、Java、操作系统、计算机网络、软件工程、人工智能等。"),
]

# ────────────────────────── 一、测试方向 ──────────────────────────
CONTENT_测试 = COMMON_EDU + [
    ("p", "主要技能：熟练 Python；软件测试；熟悉 Java、Spring Boot、MySQL；了解 Vue 3、uni-app 前端开发。"),

    ("h", "项目经历"),
    ("kv", "校园帮 —— 校园生活服务平台（微信小程序 + 管理端）", ""),
    ("p", "项目描述：面向校园生活的跑腿服务平台，含移动端微信小程序与后台管理端。本人负责订单模块后端接口、管理员登录鉴权、管理端订单与用户管理页面，并设计验收该项目的接口自动化测试。"),
    ("b", "开发模式（AI 辅助）：", "项目采用 AI 辅助开发（Vibe Coding）——我负责业务拆解、接口契约设计与测试用例设计，编码由 AI 协作产出后逐条验收。"),
    ("p", "技术栈：Java / Spring Boot / MyBatis-Plus / MySQL / JWT / Vue 3 + TypeScript / uni-app / Python + Requests + Pytest + Allure"),
    ("b", "接口自动化（主要产出）：", "基于 Python + Requests + Pytest 搭建订单模块接口自动化测试：封装会话级 Token 与鉴权 Header、订单构造与数据库连接 fixture；编写 9 条带断言的用例，覆盖 6 个接口的正常流程、状态流转与「逻辑删除后不可见」等边界规则。"),
    ("b", "报告与数据校验：", "接入 Allure 生成可视化报告，并把请求 / 响应报文附加到报告，失败用例可直接看报文定位；用例直连 MySQL 校验落库结果（逻辑删除后行仍在、deleted 置 1），并清理自建数据保证用例可反复执行。"),
    ("b", "后端开发：", "基于 Spring Boot + MyBatis-Plus 实现订单模块 6 个接口（列表、详情、新增、编辑、强制取消、删除），设计订单状态流转与「仅待支付、待接单可编辑」等业务规则；基于 JWT 实现管理员登录，通过拦截器统一校验 Token。"),
    ("b", "管理端开发：", "使用 Vue 3 + TypeScript + Element Plus 实现订单管理、用户管理页面（列表查询、编辑、禁用 / 启用），完成与后端接口联调。"),

    ("kv", "接口自动化测试练习项目（实训课设）", ""),
    ("p", "项目描述：实训课设中针对一个示例电商系统练习接口自动化测试的完整流程（课设练习项目，非商业项目）。"),
    ("b", "用例设计：", "覆盖注册 / 登录、商品浏览、购物车、订单支付等模块，练习等价类、边界值等用例设计方法。"),
    ("b", "框架搭建：", "基于 Python + Requests + Pytest 封装统一请求方法与登录鉴权逻辑，实现业务流程串联与回归测试。"),
    ("b", "版本管理：", "使用 Git / Gitee 管理测试代码与测试文档。"),

    ("h", "技能特长"),
    ("b", "测试基础：", "熟悉软件测试全流程，掌握等价类、边界值等用例设计方法，能独立输出测试报告。"),
    ("b", "接口测试：", "熟练使用 Python + Requests 进行接口测试，掌握 GET / POST 请求、参数传递、鉴权 Header 处理与断言验证。"),
    ("b", "自动化测试：", "熟练使用 Pytest，能实现用例组织、fixture 数据准备与清理、数据驱动与 Allure 测试报告生成。"),
    ("b", "开发能力：", "熟悉 Java + Spring Boot + MySQL，了解 Vue 3 + TypeScript，能与开发高效定位问题。"),
    ("b", "工具与版本控制：", "熟练使用 Git / Gitee 管理代码，了解 Postman、Fiddler 等接口测试辅助工具。"),

    ("h", "校园经历与荣誉"),
    ("b", "校声乐队优秀队员：", "参与校内文艺汇演、校外商演等十余场演出，负责主唱与和声部分，锻炼了抗压能力与舞台表现力。"),
    ("b", "英语四级：", "听说读写能力良好，能流利用英语进行日常交流，能快速浏览英文文档与书籍。"),

    ("h", "自我评价"),
    ("b", "测试能力扎实：", "软件工程专业大四学生；能独立完成功能测试与接口自动化测试，掌握 Pytest、Requests、Allure，已在实训项目中独立搭建订单模块接口自动化测试并落地执行。"),
    ("b", "具备开发经验：", "参与过完整的项目开发（后端接口 + 管理端页面），理解开发与联调过程，能站在开发视角定位缺陷根因，而不只停留在现象描述。"),
]

# ─────────────────────── 二、AI 应用开发方向 ───────────────────────
CONTENT_AI = COMMON_EDU + [
    ("p", "方向：Python 后端 + 大模型应用开发（RAG / Agent）"),

    ("h", "项目经历"),
    ("kv", "企业知识库 RAG 问答系统（HR 劳动关系合规助手）", ""),
    ("p", "技术栈：Python / LangChain / ChromaDB / bge-small-zh-v1.5 / BM25(jieba) / CrossEncoder 重排 / FastAPI"),
    ("b", "混合检索 + 重排，用消融实验量化部件分工：", "向量（bge-small-zh-v1.5）与 BM25（jieba 分词）经 RRF 融合召回，再交 bge-reranker-base 逐对重排取 Top-3；四种策略做成可切换的 mode 跑消融 —— BM25 提升召回（hit@10 0.968 → 1.000），Reranker 提升排序（hit@1 0.677 → 0.742、MRR 0.797 → 0.851），而 MMR 在本语料上是负收益（hit@10 → 0.839），故不设为默认。"),
    ("b", "零 API 成本的检索评估体系：", "只评召回不评生成（回答质量需 LLM 当裁判，贵且不可复现）。31 条用例 × 6 个指标（hit@1/@3/@k、Recall@k、MRR、nDCG@k）；判据用「来源文件 + 原文子串」而非 chunk id —— id 随切分参数变化，绑死它就无法对切分策略做 A/B。配套自检脚本先证明每条判据逐字成立，首次运行即抓出 6 条错判据。"),
    ("b", "多知识库与语料工程：", "面向中小企业 HR 的「劳动关系合规助手」——三个互相隔离的知识库（ChromaDB collection）共 1774 块：企业员工手册、本人 4 个项目文档（602 块）、劳动法律法规 7 部（1127 块）。法条单独按「第X条」边界切分，元数据落到 {法律, 版本, 编, 章, 条}，引用可精确到具体条款；按固定字数切会把法条从中间劈开、编号与正文分家。"),
    ("b", "用评估定位并修复一个生产级缺陷：", "BM25Retriever 未传 metadatas，导致 BM25 召回的块经 RRF 融合后丢失 source、接口 sources 返回 null；同一疏漏还让评估把 BM25 的贡献系统性算成失败。另含上传幂等（块 id 取内容 sha1，重复上传 added=0）与中文 BM25 分词坑（不传 jieba 分词则整句成一个 token）。"),

    ("kv", "手写 AI Agent（独立开发，不使用 LangChain）", ""),
    ("p", "技术栈：Python / OpenAI 兼容接口（DeepSeek）/ MCP 2.3 / FastAPI"),
    ("b", "Agent Loop 手写实现：", "完整实现「模型决策 → 解析 tool_calls → 本地执行 → 结果回灌 → 再决策」的多轮循环，含工具注册、参数校验、轮次上限，并覆盖超时 / 参数错 / 模型编造工具名三条错误线。"),
    ("b", "MCP Server：", "基于 Model Context Protocol 自建 Server（7 个工具 + 1 个资源模板：只读资料做成资源、需要模型主动执行的动作才做成工具）；用 Path.resolve() + is_relative_to() 做路径层级白名单防越界 —— 而非字符串前缀匹配，后者会被 docs_bak 这类兄弟目录绕过。"),
    ("b", "Agentic RAG（跨进程）：", "将已有 RAG 服务封装为 Agent 的一个工具，走 HTTP 跨进程调用而非同进程 import，避免把数 GB 的 embedding 依赖装进 Agent 项目。"),

    ("kv", "头条新闻全栈资讯平台（独立开发）", ""),
    ("p", "技术栈：FastAPI / SQLAlchemy 2.0（异步）/ MySQL / Redis / Vue 3 / Pinia"),
    ("b", "后端：", "设计 6 张数据库表、完成 18 个 RESTful API；分层架构（models → schemas → crud → routers）；Token 认证 + FastAPI Depends 依赖注入全局鉴权；删除 / 修改接口强制携带 user_id 条件防越权。"),
    ("b", "Redis 与 AI 集成：", "自建 @cached(ttl, key_builder) 装饰器接入 Cache-Aside，命中响应 ~100ms → ~2ms；浏览量改 Redis INCR 原子自增 + Write-Behind 批量回写（Lua 脚本原子取回 + 清空，落库成功才销登记）；AI 走后端代理（API Key 不外泄）+ SSE 流式转发。"),

    ("p", "智能分拣系统（AI 视觉分拣）｜团队 5 人 · 本人负责核心链路：发现课程验证集自第 1 轮即 100%（纯颜色最近邻也达 90.2%），判定样本泄漏并弃用；自建「留一法」真值判据（精度 99.7%）把总判断 71.2% 的失分定位到单一半熟类，再经剂量实验得出每视角 1~3 张即收敛 —— 考场半熟识别率 46.4% → 100%、总判断 71.2% → 98.5%。"),

    ("p", "附加：校园跑腿平台（Java 全栈 · AI 辅助开发 Vibe Coding，我负责业务与框架设计，编码由 AI 产出后验收，含 Spring AI 助手）：Spring Boot 3 三层架构后端（66 个类 / 15 张表）+ JWT 鉴权 + Vue 3 管理端 + uni-app 移动端；并基于 Spring AI 实现 Advisor 拦截链（日志 / 敏感词 / 自定义 Prompt 改写）与流式 Flux 调用。"),

    ("h", "技能特长"),
    ("b", "大模型应用：", "RAG 全流程（分段 / Embedding / 向量检索 / 重排）；Hybrid Search（向量 + BM25 + RRF）；Function Calling 与 Agent Loop 手写实现；MCP Server 开发；Prompt 工程；DeepSeek / OpenAI 兼容接口。"),
    ("b", "Python 后端：", "FastAPI、SQLAlchemy 2.0 异步 ORM、Pydantic v2、asyncio、SSE 流式响应、httpx。"),
    ("b", "数据与中间件：", "MySQL、Redis（缓存 / 原子计数器 / 连接池）、ChromaDB 向量库。"),
    ("b", "测试与其他：", "Pytest 接口自动化（数据驱动 + Allure 报告）；Git、Linux、阿里云部署；Java / Spring Boot 有完整项目经验。"),

    ("h", "自我评价"),
    ("b", "能落地且重根因：", "RAG 检索链路、Agent 循环与工具调用均为手写实现，能讲清每一轮模型往返做了什么、失败路径怎么兜；习惯用「副作用是否发生」「换随机种子重跑」这类硬判据验证结论，并据此主动否决过两项统计上不可区分的优化方案。"),
]

# ───────────────────────── 三、Java 后端方向 ─────────────────────────
CONTENT_JAVA = COMMON_EDU + [
    ("p", "主要技能：Java / Spring Boot / MyBatis-Plus / MySQL / Redis / JWT；同时熟练 Python 后端（FastAPI）与大模型应用（RAG / Agent）。"),

    ("h", "项目经历"),
    ("kv", "校园帮 —— 校园生活服务平台（Spring Boot 全栈）", ""),
    ("p", "项目描述：面向校园生活的跑腿服务平台，覆盖后端服务、后台管理端与微信小程序移动端。"),
    ("p", "技术栈：Java / Spring Boot 3.1.6 / MyBatis-Plus 3.5.3 / MySQL / Redis / JWT（jjwt）/ Druid / Vue 3 + TypeScript / uni-app"),
    ("b", "后端架构：", "Spring Boot 三层分层架构（controller / service / mapper / entity / dto / config / common / exception），66 个 Java 类、15 张数据库表（用户 / 订单 / 申诉 / 消息 / 通知 / 钱包流水 / 信用记录 / RBAC 权限体系）。"),
    ("b", "订单模块与鉴权：", "实现订单 6 个接口（列表、详情、新增、编辑、强制取消、逻辑删除），设计订单状态流转与「仅待支付、待接单可编辑」等业务规则；基于 JWT（jjwt）实现管理员登录，通过 HandlerInterceptor 拦截器统一校验 Token，密码用 BCrypt 加密存储。"),
    ("b", "管理端、统计与接口测试：", "实现订单统计 / 趋势等聚合查询接口供 Vue 3 + Element Plus 管理端展示，完成订单与用户管理页面联调；并独立搭建该项目订单模块的 Pytest + Requests + Allure 接口自动化测试（9 条带断言用例 + 报文附加 + MySQL 落库校验）。"),
    ("b", "依赖选型的工程判断：", "密码加密只引 spring-security-crypto 独立小包而不引 spring-boot-starter-security（后者会带出整条安全过滤链，导致所有接口突然要求 Basic 认证）；老师文档指定的 Druid 1.0.9 未实现 JDBC 4.1 的 getObject(int, Class)，MyBatis 读 LocalDateTime 字段会抛 SQLFeatureNotSupportedException，升级至 1.2.20 解决。"),

    ("kv", "校园百事通 —— 基于 Spring AI 的领域智能助手（独立开发）", ""),
    ("p", "技术栈：Spring Boot / Spring AI / DeepSeek API / Maven / JDK 17"),
    ("b", "Advisor 拦截链：", "通过 .defaultAdvisors(...) 挂载三个 Advisor，执行顺序由 getOrder() 决定：日志（SimpleLoggerAdvisor）、敏感词拦截（SafeGuardAdvisor，命中直接返回兜底语、不请求模型）、自定义 BaseAdvisor（用 PromptTemplate 把用户问题改写为「先给结论、再分点解释」的结构）。"),
    ("b", "三种调用模式：", "同步 call()、流式 Flux<String>、PromptTemplate 双占位符；System 人设抽到 resources/prompt/ 下与代码解耦。"),
    ("b", "踩坑与修复：", "讲义示例的 before() 整体替换 Prompt 会丢掉 System 消息导致人设失效 → 改为只重写用户消息、用 prompt.mutate() 保留 System；.system(Resource) 单参重载内部用 Charset.defaultCharset()，中文 Windows 下是 GBK 会把 UTF-8 的 txt 读成乱码 → 显式指定 StandardCharsets.UTF_8。"),

    ("kv", "头条新闻全栈资讯平台（独立开发 · Python 后端）", ""),
    ("p", "技术栈：FastAPI / SQLAlchemy 2.0（异步）/ MySQL / Redis / Vue 3 / Pinia"),
    ("b", "后端：", "设计 6 张表、完成 18 个 RESTful API；Token 认证 + Depends 依赖注入全局鉴权；删除 / 修改接口强制携带 user_id 条件防越权。"),
    ("b", "缓存与写入优化：", "自建 @cached(ttl, key_builder) 装饰器接入 Cache-Aside 缓存，命中响应从 ~100ms 降至 ~2ms；浏览量写入改为 Redis INCR 原子自增 + Write-Behind 定时批量回写 MySQL，回写用 Lua 脚本实现「原子取回 + 清空」，落库成功才销登记。"),

    ("p", "智能分拣系统（AI 视觉分拣）｜团队 5 人 · 本人负责核心链路：自建「留一法」真值判据（精度 99.7%）定位训练集覆盖缺口，考场半熟识别率 46.4% → 100%、总判断 71.2% → 98.5%，正式考试 201/202 = 99.50%。"),

    ("h", "技能特长"),
    ("b", "Java 后端：", "Spring Boot 3.x、MyBatis-Plus、Maven、JWT 鉴权、BCrypt、Druid 连接池、AOP、参数校验、全局异常处理。"),
    ("b", "数据库与中间件：", "MySQL（建表 / 索引 / 外键 / 多表关联 / 逻辑删除）、Redis（缓存 / 原子计数器 / 连接池）、WebSocket。"),
    ("b", "Spring AI / Python：", "Spring AI（Advisor 拦截链、PromptTemplate、call() 与流式 Flux、自定义 BaseAdvisor）；FastAPI、SQLAlchemy 异步 ORM、RAG、LangChain、ChromaDB、Function Calling / Agent。"),
    ("b", "前端与其他：", "Vue 3 + TypeScript + Element Plus 管理端、uni-app 移动端；Pytest 接口自动化（Allure 报告）、Git、Linux、阿里云部署。"),

    ("h", "自我评价"),
    ("b", "技术面较宽且有项目支撑：", "Java 后端、Python 后端、大模型应用三条线都独立完成过完整项目；注重根因 —— 例如从数据库驱动版本差异解释 MyBatis 抛出的异常，而不是换个依赖碰运气。"),
]

# 每套的第三个值是【行距倍数】—— 单页容量的主要旋钮。
# 内容多的那两套调小一点，尽量压在 1 页内（双选会递纸，1 页最稳）。
VARIANTS = {
    "测试": ("求职意向：软件测试 / 接口自动化测试　|　期望城市：成都　|　到岗时间：一个月内", CONTENT_测试, 1.05),
    "AI应用开发": ("求职意向：AI 应用开发 / 大模型应用开发（实习）　|　期望城市：成都　|　到岗时间：一个月内", CONTENT_AI, 0.88),
    "Java后端": ("求职意向：Java 后端开发 / 全栈开发（实习）　|　期望城市：成都　|　到岗时间：一个月内", CONTENT_JAVA, 0.92),
}

# ────────────────────────────── DOCX ──────────────────────────────
DARK = (0x1F, 0x38, 0x64)
GRAY = (0x59, 0x59, 0x59)
LS = 1.05   # 默认行距（单套生成时会被 VARIANTS 里的值覆盖）


def build_docx(line1, content, docx_path, ls=LS):
    global LS
    LS = ls
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    s = doc.styles["Normal"]
    s.font.name = "宋体"; s.font.size = Pt(10)
    s.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    s.paragraph_format.space_before = Pt(0)
    s.paragraph_format.space_after = Pt(0)
    s.paragraph_format.line_spacing = LS
    for sec in doc.sections:
        sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)   # A4
        sec.top_margin = Cm(1.0); sec.bottom_margin = Cm(0.8)
        sec.left_margin = Cm(1.5); sec.right_margin = Cm(1.5)

    def F(run, size=10, bold=False, font="宋体", color=None):
        run.font.name = font; run.font.size = Pt(size); run.font.bold = bold
        run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        if color:
            run.font.color.rgb = RGBColor(*color)
        return run

    def P(space_after=0, space_before=0, indent=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.line_spacing = LS
        if indent is not None:
            p.paragraph_format.left_indent = Cm(indent)
        return p

    def heading(text):
        p = P(space_after=1.2, space_before=2.5)
        F(p.add_run(text), 11, True, "黑体", DARK)
        pPr = p._p.get_or_add_pPr()
        bd = OxmlElement("w:pBdr")
        bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "8")
        bot.set(qn("w:space"), "2"); bot.set(qn("w:color"), "1F3864")
        bd.append(bot); pPr.append(bd)

    # 抬头
    p = P(space_after=0)
    p.paragraph_format.line_spacing = 1.0
    F(p.add_run(NAME), 19, True, "黑体", DARK)
    p = P(space_after=0)
    F(p.add_run(line1), 11, True, "宋体", DARK)
    p = P(space_after=0)
    F(p.add_run(LINE2), 9.5, False, "宋体", GRAY)
    p = P(space_after=4)
    F(p.add_run(LINE3), 9.5, False, "宋体", GRAY)

    for blk in content:
        if blk[0] == "h":
            heading(blk[1])
        elif blk[0] == "kv":
            p = P(space_after=1, space_before=3)
            F(p.add_run(blk[1]), 10, True)
            if blk[2]:
                F(p.add_run("　　" + blk[2]), 9.5, False, "宋体", GRAY)
        elif blk[0] == "p":
            p = P(space_after=1, indent=0.38)
            F(p.add_run(blk[1]), 10)
        else:
            p = P(space_after=1, indent=0.38)
            p.paragraph_format.first_line_indent = Cm(-0.38)
            F(p.add_run("● "), 7, False, "宋体", DARK)
            F(p.add_run(blk[1]), 10, True)
            F(p.add_run(blk[2]), 10)

    doc.save(docx_path)
    return os.path.getsize(docx_path)


# ────────────────────────────── PDF ──────────────────────────────
def build_pdf(line1, content, pdf_path, ls=LS):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                    HRFlowable)

    pdfmetrics.registerFont(TTFont("SimSun", r"C:\Windows\Fonts\simsun.ttc", subfontIndex=0))
    pdfmetrics.registerFont(TTFont("SimHei", r"C:\Windows\Fonts\simhei.ttf"))

    dark = colors.HexColor("#1F3864")
    gray = colors.HexColor("#595959")

    body = ParagraphStyle("body", fontName="SimSun", fontSize=10, leading=12.0 * ls,
                          textColor=colors.black, firstLineIndent=10.8,
                          spaceAfter=0.5)
    bullet = ParagraphStyle("bullet", parent=body, firstLineIndent=-10.8,
                            leftIndent=10.8)
    head = ParagraphStyle("head", fontName="SimHei", fontSize=11, leading=13.5,
                          textColor=dark, spaceBefore=3, spaceAfter=0.5)
    kv = ParagraphStyle("kv", parent=body, firstLineIndent=0, spaceBefore=2.5)
    hdr = ParagraphStyle("hdr", parent=body, firstLineIndent=0, fontSize=19,
                         leading=25, spaceAfter=0)
    intent = ParagraphStyle("intent", parent=body, firstLineIndent=0, fontSize=11,
                            leading=14, textColor=dark, fontName="SimHei")
    sub = ParagraphStyle("sub", parent=body, firstLineIndent=0, fontSize=9.5,
                         leading=11.5, textColor=gray)

    doc = SimpleDocTemplate(pdf_path, pagesize=A4,
                            leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                            topMargin=1.1 * cm, bottomMargin=0.9 * cm,
                            title=os.path.basename(pdf_path), author=NAME)
    story = []

    story.append(Paragraph(NAME, hdr))
    story.append(Paragraph(line1, intent))
    story.append(Paragraph(LINE2, sub))
    story.append(Paragraph(LINE3, sub))
    story.append(Spacer(1, 2))

    for blk in content:
        if blk[0] == "h":
            story.append(Paragraph(blk[1], head))
            story.append(HRFlowable(width="100%", thickness=0.7, color=dark,
                                    spaceBefore=0, spaceAfter=2.5))
        elif blk[0] == "kv":
            txt = "<b>%s</b>" % blk[1]
            if blk[2]:
                txt += '<font size="9.5" color="#595959">　　%s</font>' % blk[2]
            story.append(Paragraph(txt, kv))
        elif blk[0] == "p":
            story.append(Paragraph(blk[1], body))
        else:
            story.append(Paragraph(
                '<font size="7" color="#1F3864">● </font><b>%s</b>%s' % (blk[1], blk[2]),
                bullet))

    doc.build(story)
    return os.path.getsize(pdf_path)


def docx_to_pdf(docx_path, pdf_path):
    """用本机 WPS（注册为 Word.Application）把 docx 导成 PDF——docx 是唯一源，
    PDF 只是它的渲染结果，两者不会走样。"""
    import win32com.client as wc
    app = wc.DispatchEx("Word.Application")
    app.Visible = False
    try:
        app.DisplayAlerts = 0
    except Exception:
        pass
    try:
        d = app.Documents.Open(docx_path, ReadOnly=True)
        try:
            d.ExportAsFixedFormat(OutputFileName=pdf_path, ExportFormat=17)
        except Exception:
            d.SaveAs2(pdf_path, FileFormat=17)
        d.Close(False)
    finally:
        try:
            app.Quit()
        except Exception:
            pass
    return os.path.getsize(pdf_path)


def main():
    want = sys.argv[1:] or list(VARIANTS)
    for name in want:
        if name not in VARIANTS:
            print("跳过（没有这一套）：%s    可选：%s" % (name, " / ".join(VARIANTS)))
            continue
        line1, content, ls = VARIANTS[name]
        base = "%s_%s" % (BASE, name)
        docx_path = os.path.join(OUT_DIR, base + ".docx")
        pdf_path = os.path.join(OUT_DIR, base + ".pdf")
        print("=" * 64)
        print("【%s】%s   (行距 %.2f)" % (name, line1, ls))
        print("  docx: %s  (%d B)" % (docx_path, build_docx(line1, content, docx_path, ls)))
        # 三套连续转换复用一个 WPS 实例容易出 COM 问题，每套独立 try
        try:
            print("  pdf : %s  (%d B)  [WPS 渲染]" % (pdf_path, docx_to_pdf(docx_path, pdf_path)))
        except Exception as e:
            print("  WPS 转换失败（%s: %s），改用内置排版" % (type(e).__name__, e))
            print("  pdf : %s  (%d B)  [内置排版]" % (pdf_path, build_pdf(line1, content, pdf_path, ls)))


if __name__ == "__main__":
    main()
